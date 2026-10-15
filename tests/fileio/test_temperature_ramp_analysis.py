"""Tests for TemperatureRampAnalysis transform.

Fixture file: test_temp_ramp.txt  (one temperature-ramp step, -40 °C → ~200 °C)
"""

import os

import numpy as np
import pytest

from piblin.data import ExperimentSet, OneDimensionalCompositeDataset
from hermes_rheo.file_readers.trios_rheo_txt import TriosRheoReader
from hermes_rheo.transforms.owchirp_analysis import OWChirpAnalysis
from hermes_rheo.transforms.temperature_ramp_analysis import (
    TemperatureRampAnalysis,
    TemperatureRampResult,
)

_USE_TEMPS = [0.0, 50.0, 100.0, 150.0]


def _run(datadir, steps_to_merge=None, use_temps=None):
    filepath = datadir.join("test_temp_ramp.txt")
    mset = TriosRheoReader().data_from_filepath(filepath=filepath, create_composite_datasets=True)
    exp = ExperimentSet(mset.measurements)
    transform = TemperatureRampAnalysis(
        steps_to_merge=steps_to_merge,
        use_temps=use_temps if use_temps is not None else _USE_TEMPS,
    )
    result_exp = (OWChirpAnalysis() + transform)(exp)
    return transform, result_exp


# ---------------------------------------------------------------------------
# Result types and attributes
# ---------------------------------------------------------------------------


def test_result_types(datadir):
    """transform.result is a TemperatureRampResult with the correct attribute types."""
    transform, result_exp = _run(datadir, steps_to_merge=["Temperature Ramp"])
    result = transform.result
    assert isinstance(result, TemperatureRampResult)
    assert isinstance(result.dataset, OneDimensionalCompositeDataset)
    assert isinstance(result.tg, float)
    assert isinstance(result.extracted_dataset, OneDimensionalCompositeDataset)
    # Shortcut attributes are also set on the returned experiment
    assert result_exp.tg == result.tg
    assert result_exp.dataset is result.dataset


# ---------------------------------------------------------------------------
# Merged dataset
# ---------------------------------------------------------------------------


def test_dataset_has_temperature_sorted_ascending(datadir):
    """Merged dataset contains a temperature variable sorted in ascending order."""
    transform, _ = _run(datadir, steps_to_merge=["Temperature Ramp"])
    ds = transform.result.dataset
    names_lower = [n.lower() for n in ds.data_array_names]
    temp_idx = names_lower.index("temperature")
    temps = np.asarray(ds.data_arrays[temp_idx], dtype=float)
    assert len(temps) > 0
    assert np.all(temps[:-1] <= temps[1:])


# ---------------------------------------------------------------------------
# Tg detection
# ---------------------------------------------------------------------------


def test_tg_within_temperature_range(datadir):
    """Detected Tg lies within the measured temperature range."""
    transform, _ = _run(datadir, steps_to_merge=["Temperature Ramp"])
    result = transform.result
    ds = result.dataset
    names_lower = [n.lower() for n in ds.data_array_names]
    temps = np.asarray(ds.data_arrays[names_lower.index("temperature")], dtype=float)
    assert temps.min() <= result.tg <= temps.max()


# ---------------------------------------------------------------------------
# Extraction at use_temps
# ---------------------------------------------------------------------------


def test_extracted_dataset_size(datadir):
    """Extracted dataset has one point per requested temperature."""
    transform, _ = _run(datadir, steps_to_merge=["Temperature Ramp"], use_temps=_USE_TEMPS)
    ext = transform.result.extracted_dataset
    assert len(ext.data_arrays[0]) == len(_USE_TEMPS)


# ---------------------------------------------------------------------------
# No matching steps (edge case)
# ---------------------------------------------------------------------------


def test_no_matching_steps_returns_none(datadir):
    """When no steps match, tg and both datasets are None."""
    transform, _ = _run(datadir, steps_to_merge=["NonExistentStep"], use_temps=[25.0])
    result = transform.result
    assert result.tg is None
    assert result.dataset is None
    assert result.extracted_dataset is None


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


def test_to_excel(datadir, tmp_path):
    """to_excel() creates a workbook with 'Temperature Series' and 'Extracted Values' sheets."""
    import openpyxl

    transform, _ = _run(datadir, steps_to_merge=["Temperature Ramp"], use_temps=[0.0, 100.0])
    out = str(tmp_path / "out.xlsx")
    transform.result.to_excel(out)
    assert os.path.exists(out)
    wb = openpyxl.load_workbook(out)
    assert "Temperature Series" in wb.sheetnames
    assert "Extracted Values" in wb.sheetnames
